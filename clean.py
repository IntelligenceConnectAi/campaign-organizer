import streamlit as st
import pandas as pd
import csv
import os
import io
import zipfile
import math
import warnings
from docx import Document
from docx.shared import Pt, RGBColor
warnings.filterwarnings("ignore")

# ── CONFIG ───────────────────────────────────────────────────────────────────
PHONE_GROUPS = [
    ("Phone 1", "Phone 1 Type", "Email 1"),
    ("Phone 2", "Phone 2 Type", "Email 2"),
    ("Phone 3", "Phone 3 Type", "Email 3"),
    ("Phone 4", "Phone 4 Type", "Email 4"),
    ("Phone 5", "Phone 5 Type", "Email 5"),
]

DNC_COLS = ["Phone 1 DNC", "Phone 2 DNC", "Phone 3 DNC", "Phone 4 DNC", "Phone 5 DNC"]

MLS_STATUSES = [
    "ACTIVE", "ACTIVE UNDER CONTRACT", "CANCELED", "CANCELLED", "COMING SOON",
    "CONTINGENT", "DELETED", "EXPIRED", "FAIL", "FAILED",
    "PENDING", "REMOVED", "WITHDRAWN"
]

MONTHS = ["JAN","FEB","MAR","APR","MAY","JUN",
          "JUL","AUG","SEP","OCT","NOV","DEC"]

MONTH_FULL = {
    "JAN":"January","FEB":"February","MAR":"March","APR":"April",
    "MAY":"May","JUN":"June","JUL":"July","AUG":"August",
    "SEP":"September","OCT":"October","NOV":"November","DEC":"December"
}

DEAL_TYPES = [
    "LUXURY LAND","NON-LUXURY","LAND LUXURY SFH","NON-LUXURY SFH",
    "TEARDOWN","INFILL LOT","LARGE ACREAGE","BUILDER LOT",
]

# Business Leads use a single Phone / single Email column (e.g. Google Maps
# scrapes) instead of the grouped Phone 1..5 format used by Sales/MLS.
BIZ_PHONE_COL = "Phone"
BIZ_EMAIL_COL = "Email"

TMP_MERGED       = "/tmp/merged.csv"
TMP_SALES        = "/tmp/sales.csv"
TMP_MLS          = "/tmp/mls.csv"
TMP_DIALER       = "/tmp/dialer_output.csv"
TMP_SMS          = "/tmp/sms_output.csv"
TMP_EMAIL        = "/tmp/email_output.csv"

# MLS Seller List temp outputs
TMP_MLS_DIALER   = "/tmp/mls_seller_dialer_output.csv"
TMP_MLS_SMS      = "/tmp/mls_seller_sms_output.csv"
TMP_MLS_EMAIL    = "/tmp/mls_seller_email_output.csv"

# MLS Agent List temp outputs
TMP_MLS_AGENT_DIALER = "/tmp/mls_agent_dialer_output.csv"
TMP_MLS_AGENT_SMS    = "/tmp/mls_agent_sms_output.csv"
TMP_MLS_AGENT_EMAIL  = "/tmp/mls_agent_email_output.csv"

# Business Leads temp outputs
TMP_BIZ_MERGED   = "/tmp/biz_merged.csv"
TMP_BIZ_DIALER   = "/tmp/biz_dialer_output.csv"
TMP_BIZ_SMS      = "/tmp/biz_sms_output.csv"
TMP_BIZ_EMAIL    = "/tmp/biz_email_output.csv"

# ── SESSION STATE ────────────────────────────────────────────────────────────
for key, val in {
    "processed": False,
    "sales_results": {},
    "mls_results": {},
    "total_merged": 0,
    "orig_phone_count": 0,
    "orig_email_count": 0,
    "sales_zip_buffer": None,
    "mls_seller_zip_buffer": None,
    "mls_agent_zip_buffer": None,
    "sales_zip_name": "",
    "mls_seller_zip_name": "",
    "mls_agent_zip_name": "",
    "report_buffer": None,
    # Business Leads
    "biz_processed": False,
    "biz_results": {},
    "biz_total_merged": 0,
    "biz_orig_phone_count": 0,
    "biz_orig_email_count": 0,
    "biz_zip_buffer": None,
    "biz_zip_name": "",
    "biz_report_buffer": None,
    # Navigation
    "current_page": "List Cleaner",
}.items():
    if key not in st.session_state:
        st.session_state[key] = val

# ── HELPERS ──────────────────────────────────────────────────────────────────
def get_val(row, col):
    if col not in row.index:
        return ""
    v = str(row[col]).strip()
    return "" if v.lower() == "nan" else v

def clean_phone(val):
    if not val:
        return ""
    try:
        return str(int(float(val)))
    except:
        return val

def clean_email(val):
    """Trims/lowercases the email and blanks it out if it doesn't look like a
    valid address (no '@', or no '.' in the domain part), so malformed or
    junk emails get dropped the same way blank ones do."""
    if not val:
        return ""
    val = str(val).strip()
    if not val or val.lower() == "nan":
        return ""
    if "@" not in val:
        return ""
    local, _, domain = val.partition("@")
    if not local or "." not in domain:
        return ""
    return val.lower()

def make_campaign(channel, month, year, state, deal, mls=False, mls_type=None, wk=None):
    if mls:
        mls_part = f" - MLS - {mls_type}" if mls_type else " - MLS"
    else:
        mls_part = ""
    wk_part  = f" - WK{wk}" if wk else ""
    return f"{channel}{mls_part}{wk_part} - {month} - {year} - {state} - {deal}"

def calc_week_ranges(total, n_splits):
    chunk = math.ceil(total / n_splits)
    ranges = []
    for i in range(n_splits):
        start = i * chunk + 1
        end   = min((i + 1) * chunk, total)
        ranges.append(f"WK{i+1}: {start:,} – {end:,} ({end-start+1:,} contacts)")
    return ranges

# ── MERGE + SPLIT SALES / MLS ────────────────────────────────────────────────
def merge_and_split(uploaded_files):
    total_rows   = 0
    orig_phones  = 0
    orig_emails  = 0
    header_written = False
    phone_cols = [p for p, t, e in PHONE_GROUPS]
    email_cols = [e for p, t, e in PHONE_GROUPS]

    with open(TMP_MERGED, "w", newline="", encoding="utf-8") as out_f:
        for uf in uploaded_files:
            df = pd.read_excel(uf, dtype=str)
            df.drop(columns=[c for c in DNC_COLS if c in df.columns], inplace=True)
            for col in phone_cols:
                if col in df.columns:
                    orig_phones += df[col].apply(
                        lambda x: 1 if pd.notna(x) and str(x).strip() not in ("","nan") else 0).sum()
            for col in email_cols:
                if col in df.columns:
                    orig_emails += df[col].apply(
                        lambda x: 1 if pd.notna(x) and str(x).strip() not in ("","nan") else 0).sum()
            df.to_csv(out_f, index=False, header=not header_written)
            total_rows += len(df)
            header_written = True
            del df

    # Sales file = full merged file as-is
    df = pd.read_csv(TMP_MERGED, dtype=str)
    df.to_csv(TMP_SALES, index=False)

    # MLS file = blank rows (→ "Off Market") + known status rows
    if "MLS Status" in df.columns:
        blank_mask = df["MLS Status"].apply(lambda x: str(x).strip() in ("", "nan"))
        known_mask = df["MLS Status"].apply(lambda x: str(x).strip().upper() in MLS_STATUSES)

        mls_known        = df[known_mask].copy()
        mls_blanks       = df[blank_mask].copy()
        mls_blanks["MLS Status"] = "Off Market"
        mls_df = pd.concat([mls_known, mls_blanks], ignore_index=True)
    else:
        mls_df = pd.DataFrame(columns=df.columns)

    mls_df.to_csv(TMP_MLS, index=False)
    del df, mls_df

    return total_rows, int(orig_phones), int(orig_emails)

# ── MERGE (Business Leads: single Phone / Email columns, xlsx + csv) ────────
def merge_biz_to_disk(uploaded_files):
    total_rows   = 0
    orig_phones  = 0
    orig_emails  = 0
    header_written = False

    with open(TMP_BIZ_MERGED, "w", newline="", encoding="utf-8") as out_f:
        for uf in uploaded_files:
            if uf.name.lower().endswith(".csv"):
                df = pd.read_csv(uf, dtype=str)
            else:
                df = pd.read_excel(uf, dtype=str)
            df.drop(columns=[c for c in DNC_COLS if c in df.columns], inplace=True)
            if BIZ_PHONE_COL in df.columns:
                orig_phones += df[BIZ_PHONE_COL].apply(
                    lambda x: 1 if pd.notna(x) and str(x).strip() not in ("","nan") else 0).sum()
            if BIZ_EMAIL_COL in df.columns:
                orig_emails += df[BIZ_EMAIL_COL].apply(
                    lambda x: 1 if pd.notna(x) and str(x).strip() not in ("","nan") else 0).sum()
            df.to_csv(out_f, index=False, header=not header_written)
            total_rows += len(df)
            header_written = True
            del df

    return total_rows, int(orig_phones), int(orig_emails)

# ── GENERIC PROCESS FUNCTIONS (Sales/MLS: grouped Phone 1..5 columns) ───────
def process_dialer_file(src_csv, out_csv, campaign_name):
    df = pd.read_csv(src_csv, dtype=str)
    all_phone_cols = [p for p, t, e in PHONE_GROUPS] + [t for p, t, e in PHONE_GROUPS]
    other_cols     = [c for c in df.columns if c not in all_phone_cols]
    active_groups  = [(p, t) for p, t, e in PHONE_GROUPS if p in df.columns]
    output_cols    = ["Campaign Name", "Phone Number", "Phone Type"] + other_cols
    total = 0
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=output_cols)
        writer.writeheader()
        for _, row in df.iterrows():
            for phone_col, type_col in active_groups:
                phone_val = clean_phone(get_val(row, phone_col))
                if not phone_val:
                    continue
                new_row = {"Campaign Name": campaign_name,
                           "Phone Number": phone_val,
                           "Phone Type":   get_val(row, type_col)}
                for col in other_cols:
                    new_row[col] = get_val(row, col)
                writer.writerow(new_row)
                total += 1
    del df
    return total

def process_sms_file(src_csv, out_csv, campaign_name):
    df = pd.read_csv(src_csv, dtype=str)
    all_phone_email_cols = ([p for p, t, e in PHONE_GROUPS] +
                            [t for p, t, e in PHONE_GROUPS] +
                            [e for p, t, e in PHONE_GROUPS if e in df.columns])
    other_cols    = [c for c in df.columns if c not in all_phone_email_cols]
    active_groups = [(p, t, e) for p, t, e in PHONE_GROUPS if p in df.columns or e in df.columns]
    output_cols   = ["Campaign Name", "Phone Number", "Phone Type", "Email"] + other_cols
    total = 0
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=output_cols)
        writer.writeheader()
        for _, row in df.iterrows():
            for phone_col, type_col, email_col in active_groups:
                phone_val  = clean_phone(get_val(row, phone_col))
                phone_type = get_val(row, type_col)
                email_val  = clean_email(get_val(row, email_col))
                if phone_type.strip().lower() == "landline":
                    phone_val = ""
                if not phone_val and not email_val:
                    continue
                new_row = {"Campaign Name": campaign_name,
                           "Phone Number": phone_val,
                           "Phone Type":   phone_type,
                           "Email":        email_val}
                for col in other_cols:
                    new_row[col] = get_val(row, col)
                writer.writerow(new_row)
                total += 1
    del df
    return total

def process_email_file(src_csv, out_csv, campaign_name):
    df = pd.read_csv(src_csv, dtype=str)
    all_phone_email_cols = ([p for p, t, e in PHONE_GROUPS] +
                            [t for p, t, e in PHONE_GROUPS] +
                            [e for p, t, e in PHONE_GROUPS if e in df.columns])
    other_cols    = [c for c in df.columns if c not in all_phone_email_cols]
    active_groups = [(p, t, e) for p, t, e in PHONE_GROUPS if p in df.columns or e in df.columns]
    output_cols   = ["Campaign Name", "Phone Number", "Phone Type", "Email"] + other_cols
    total = 0
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=output_cols)
        writer.writeheader()
        for _, row in df.iterrows():
            for phone_col, type_col, email_col in active_groups:
                email_val = clean_email(get_val(row, email_col))
                if not email_val:
                    continue
                phone_val  = clean_phone(get_val(row, phone_col))
                phone_type = get_val(row, type_col)
                if phone_type.strip().lower() == "landline":
                    phone_val = ""
                new_row = {"Campaign Name": campaign_name,
                           "Phone Number": phone_val,
                           "Phone Type":   phone_type,
                           "Email":        email_val}
                for col in other_cols:
                    new_row[col] = get_val(row, col)
                writer.writerow(new_row)
                total += 1
    del df
    return total

# ── BUSINESS LEADS PROCESS FUNCTIONS (single Phone / Email columns) ─────────
def process_biz_dialer_file(src_csv, out_csv, campaign_name):
    df = pd.read_csv(src_csv, dtype=str)
    other_cols  = [c for c in df.columns if c != BIZ_PHONE_COL]
    output_cols = ["Campaign Name", "Phone Number"] + other_cols
    total = 0
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=output_cols)
        writer.writeheader()
        for _, row in df.iterrows():
            phone_val = clean_phone(get_val(row, BIZ_PHONE_COL))
            if not phone_val:
                continue
            new_row = {"Campaign Name": campaign_name, "Phone Number": phone_val}
            for col in other_cols:
                new_row[col] = get_val(row, col)
            writer.writerow(new_row)
            total += 1
    del df
    return total

def process_biz_sms_file(src_csv, out_csv, campaign_name):
    df = pd.read_csv(src_csv, dtype=str)
    other_cols  = [c for c in df.columns if c not in (BIZ_PHONE_COL, BIZ_EMAIL_COL)]
    output_cols = ["Campaign Name", "Phone Number", "Email"] + other_cols
    total = 0
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=output_cols)
        writer.writeheader()
        for _, row in df.iterrows():
            phone_val = clean_phone(get_val(row, BIZ_PHONE_COL))
            email_val = clean_email(get_val(row, BIZ_EMAIL_COL))
            if not phone_val and not email_val:
                continue
            new_row = {"Campaign Name": campaign_name, "Phone Number": phone_val, "Email": email_val}
            for col in other_cols:
                new_row[col] = get_val(row, col)
            writer.writerow(new_row)
            total += 1
    del df
    return total

def process_biz_email_file(src_csv, out_csv, campaign_name):
    df = pd.read_csv(src_csv, dtype=str)
    other_cols  = [c for c in df.columns if c not in (BIZ_PHONE_COL, BIZ_EMAIL_COL)]
    output_cols = ["Campaign Name", "Phone Number", "Email"] + other_cols
    total = 0
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=output_cols)
        writer.writeheader()
        for _, row in df.iterrows():
            email_val = clean_email(get_val(row, BIZ_EMAIL_COL))
            if not email_val:
                continue
            phone_val = clean_phone(get_val(row, BIZ_PHONE_COL))
            new_row = {"Campaign Name": campaign_name, "Phone Number": phone_val, "Email": email_val}
            for col in other_cols:
                new_row[col] = get_val(row, col)
            writer.writerow(new_row)
            total += 1
    del df
    return total

# ── SPLIT (shared by Sales/MLS and Business Leads) ──────────────────────────
def split_csv(tmp_path, channel, month, year, state, deal, n_splits, mls=False, mls_type=None):
    df = pd.read_csv(tmp_path, dtype=str)
    total      = len(df)
    chunk_size = math.ceil(total / n_splits)
    parts      = []
    for i in range(n_splits):
        chunk = df.iloc[i * chunk_size:(i + 1) * chunk_size].copy()
        if chunk.empty:
            continue
        wk_name = make_campaign(channel, month, year, state, deal, mls=mls, mls_type=mls_type, wk=i+1)
        chunk["Campaign Name"] = wk_name
        buf = io.StringIO()
        chunk.to_csv(buf, index=False)
        parts.append((wk_name, buf.getvalue().encode("utf-8")))
    del df
    return parts

# ── WORD REPORT ──────────────────────────────────────────────────────────────
def generate_report(month, year, state, deal,
                    sales_results, mls_results,
                    orig_phone_count, orig_email_count,
                    n_splits, do_split,
                    section_title="── SALES PROCESS ──"):
    doc = Document()

    def add_heading(text, level=1):
        p   = doc.add_paragraph()
        run = p.add_run(text)
        run.bold           = True
        run.font.size      = Pt(14 if level == 1 else 12)
        run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        return p

    def add_bullet(label, value):
        p         = doc.add_paragraph(style="List Bullet")
        run_label = p.add_run(f"{label}: ")
        run_label.bold = True
        p.add_run(str(value))

    def add_numbered_list(items):
        for item in items:
            p = doc.add_paragraph(style="List Number")
            p.add_run(item)

    LABELS = {"dialer": "COLD CALL", "sms": "SMS", "email": "EMAIL"}
    PHONE_LABEL = {"dialer": "Mobile and Land Lines", "sms": "Mobile Lines", "email": None}

    def add_section(title, results, orig_phones, orig_emails):
        add_heading(title, level=1)
        for key, info in results.items():
            channel_key = info["channel_key"]
            wk_ranges = calc_week_ranges(info["rows"], n_splits) if do_split and n_splits > 1 else []
            doc.add_paragraph()
            add_heading(f"{LABELS[channel_key]} Campaign Summary", level=2)
            add_bullet("Campaign Name", info["campaign"])
            if channel_key in ("dialer", "sms"):
                add_bullet(f"Total Original Count of {PHONE_LABEL[channel_key]}", f"{orig_phones:,}")
                add_bullet(f"Total Cleaned Count of {PHONE_LABEL[channel_key]}", f"{info['rows']:,}")
            else:
                add_bullet("Total Original Count of Emails", f"{orig_emails:,}")
                add_bullet("Total Cleaned Count of Emails", f"{info['rows']:,}")
            add_bullet("Week Range", " | ".join(wk_ranges) if wk_ranges else "N/A")
            add_bullet("List Type", deal)
            add_bullet("Round Progress", "[To be updated]")
            if info.get("splits"):
                p   = doc.add_paragraph()
                r   = p.add_run("WEEKLY SPLIT:")
                r.bold = True
                add_numbered_list([wk_name for wk_name, _ in info["splits"]])

    month_full = MONTH_FULL.get(month, month)
    p   = doc.add_paragraph()
    r   = p.add_run(f"Subject: Upcoming Campaign-Ready Cleaned List for {month_full} - {year}")
    r.bold          = True
    r.font.size     = Pt(13)
    doc.add_paragraph()
    doc.add_paragraph("Hi Team,")
    doc.add_paragraph(
        f"The contact list for {month_full} - {year} has been cleaned and validated. "
        "It is now organized and ready for use across the following marketing channels:"
    )
    add_heading("Details:", level=2)
    add_bullet("State(s)", state)
    add_bullet("Cities", "[To be updated]")
    add_bullet("Month", month_full)
    add_bullet("Year", year)
    add_bullet("Niche(s)", deal)

    if sales_results:
        doc.add_paragraph()
        add_section(section_title, sales_results,
                    orig_phone_count, orig_email_count)

    # MLS results are split into Seller / Agent groups for the report
    seller_results = {k: v for k, v in mls_results.items() if v.get("mls_type") == "SELLER"}
    agent_results  = {k: v for k, v in mls_results.items() if v.get("mls_type") == "AGENT"}

    if seller_results:
        doc.add_paragraph()
        add_section("── MLS SELLER LIST ──", seller_results,
                    orig_phone_count, orig_email_count)

    if agent_results:
        doc.add_paragraph()
        add_section("── MLS AGENT LIST ──", agent_results,
                    orig_phone_count, orig_email_count)

    doc.add_paragraph()
    doc.add_paragraph(
        "Please review the attached list and let me know if you have any questions or need any additional updates."
    )

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

# ══════════════════════════════════════════════════════════════════════════
# PAGE: LIST CLEANER (Sales & MLS)
# ══════════════════════════════════════════════════════════════════════════
def page_sales_mls():
    st.title("📋 List Cleaner (Sales & MLS)")
    st.divider()

    # STEP 1 — Upload
    st.subheader("Step 1 — Upload Files")
    uploaded_files = st.file_uploader(
        "📁 Upload one or more Excel files (.xlsx)",
        type=["xlsx"], accept_multiple_files=True, key="sm_upload"
    )
    if uploaded_files:
        st.success(f"✅ {len(uploaded_files)} file(s) uploaded")
    st.divider()

    # STEP 2 — Campaign Details
    st.subheader("Step 2 — Campaign Details")
    col1, col2 = st.columns(2)
    with col1:
        month = st.selectbox("📅 Month", MONTHS, key="sm_month")
        state = st.text_input("🗺️ State", placeholder="e.g. FL", key="sm_state")
    with col2:
        year  = st.text_input("📆 Year", value="2026", key="sm_year")
        deal  = st.selectbox("🏷️ Type of Deal", DEAL_TYPES, key="sm_deal")
    st.divider()

    # STEP 3 — Sales Process Options
    st.subheader("Step 3 — Sales Process")
    col_d, col_s, col_e = st.columns(3)
    with col_d:
        sales_dialer = st.checkbox("📞 Dialer", key="s_dialer")
    with col_s:
        sales_sms    = st.checkbox("💬 SMS", key="s_sms")
    with col_e:
        sales_email  = st.checkbox("📧 Email", key="s_email")
    st.divider()

    # STEP 4 — MLS Seller List
    st.subheader("Step 4 — MLS Seller List")
    col_md, col_ms, col_me = st.columns(3)
    with col_md:
        mls_seller_dialer = st.checkbox("📞 Dialer", key="m_dialer")
    with col_ms:
        mls_seller_sms    = st.checkbox("💬 SMS", key="m_sms")
    with col_me:
        mls_seller_email  = st.checkbox("📧 Email", key="m_email")
    st.divider()

    # STEP 5 — MLS Agent List (optional)
    st.subheader("Step 5 — MLS Agent List (optional)")
    col_ad, col_as, col_ae = st.columns(3)
    with col_ad:
        mls_agent_dialer = st.checkbox("📞 Dialer", key="a_dialer")
    with col_as:
        mls_agent_sms    = st.checkbox("💬 SMS", key="a_sms")
    with col_ae:
        mls_agent_email  = st.checkbox("📧 Email", key="a_email")
    st.divider()

    # STEP 6 — Split
    st.subheader("Step 6 — Split (Optional)")
    do_split = st.checkbox("🔀 Split output into multiple files", key="sm_dosplit")
    n_splits = 1
    if do_split:
        n_splits = st.number_input("How many files to split into?",
                                   min_value=2, max_value=50, value=5, step=1, key="sm_nsplits")
    st.divider()

    # PROCESS
    sales_selected      = sales_dialer or sales_sms or sales_email
    mls_seller_selected = mls_seller_dialer or mls_seller_sms or mls_seller_email
    mls_agent_selected  = mls_agent_dialer or mls_agent_sms or mls_agent_email
    mls_selected        = mls_seller_selected or mls_agent_selected
    ready = bool(uploaded_files and state.strip() and year.strip() and
                 (sales_selected or mls_selected))

    if st.button("⚙️ Process Files", use_container_width=True, type="primary",
                 disabled=not ready, key="sm_process_btn"):
        st.session_state.processed      = False
        st.session_state.sales_results  = {}
        st.session_state.mls_results    = {}

        # Merge + split Sales/MLS
        with st.spinner("Merging and splitting files... ⏳"):
            try:
                total, orig_phones, orig_emails = merge_and_split(uploaded_files)
                st.session_state.total_merged     = total
                st.session_state.orig_phone_count = orig_phones
                st.session_state.orig_email_count = orig_emails
            except Exception as e:
                st.error(f"❌ Merge error: {e}")
                st.stop()

        y = year.strip()
        s = state.strip().upper()

        # ── SALES PROCESSING ─────────────────────────────────────────────────
        SALES_MAP = {
            "dialer": (sales_dialer, TMP_DIALER,     process_dialer_file, "CC"),
            "sms":    (sales_sms,    TMP_SMS,         process_sms_file,    "SMS"),
            "email":  (sales_email,  TMP_EMAIL,       process_email_file,  "EMAIL"),
        }
        for key, (selected, out_tmp, fn, ch) in SALES_MAP.items():
            if not selected:
                continue
            campaign = make_campaign(ch, month, y, s, deal)
            with st.spinner(f"Processing Sales {key.upper()}... ⏳"):
                try:
                    rows = fn(TMP_SALES, out_tmp, campaign)
                    st.session_state.sales_results[key] = {
                        "channel": ch, "channel_key": key, "campaign": campaign,
                        "rows": rows, "tmp": out_tmp, "splits": None,
                    }
                except Exception as e:
                    st.error(f"❌ Sales {key} error: {e}")

        # ── MLS SELLER LIST PROCESSING ───────────────────────────────────────
        MLS_SELLER_MAP = {
            "dialer": (mls_seller_dialer, TMP_MLS_DIALER,  process_dialer_file, "CC"),
            "sms":    (mls_seller_sms,    TMP_MLS_SMS,      process_sms_file,    "SMS"),
            "email":  (mls_seller_email,  TMP_MLS_EMAIL,    process_email_file,  "EMAIL"),
        }
        for key, (selected, out_tmp, fn, ch) in MLS_SELLER_MAP.items():
            if not selected:
                continue
            campaign = make_campaign(ch, month, y, s, deal, mls=True, mls_type="SELLER")
            with st.spinner(f"Processing MLS Seller {key.upper()}... ⏳"):
                try:
                    rows = fn(TMP_MLS, out_tmp, campaign)
                    st.session_state.mls_results[f"seller_{key}"] = {
                        "channel": ch, "channel_key": key, "campaign": campaign,
                        "rows": rows, "tmp": out_tmp, "splits": None, "mls_type": "SELLER",
                    }
                except Exception as e:
                    st.error(f"❌ MLS Seller {key} error: {e}")

        # ── MLS AGENT LIST PROCESSING (optional) ─────────────────────────────
        MLS_AGENT_MAP = {
            "dialer": (mls_agent_dialer, TMP_MLS_AGENT_DIALER, process_dialer_file, "CC"),
            "sms":    (mls_agent_sms,    TMP_MLS_AGENT_SMS,    process_sms_file,    "SMS"),
            "email":  (mls_agent_email,  TMP_MLS_AGENT_EMAIL,  process_email_file,  "EMAIL"),
        }
        for key, (selected, out_tmp, fn, ch) in MLS_AGENT_MAP.items():
            if not selected:
                continue
            campaign = make_campaign(ch, month, y, s, deal, mls=True, mls_type="AGENT")
            with st.spinner(f"Processing MLS Agent {key.upper()}... ⏳"):
                try:
                    rows = fn(TMP_MLS, out_tmp, campaign)
                    st.session_state.mls_results[f"agent_{key}"] = {
                        "channel": ch, "channel_key": key, "campaign": campaign,
                        "rows": rows, "tmp": out_tmp, "splits": None, "mls_type": "AGENT",
                    }
                except Exception as e:
                    st.error(f"❌ MLS Agent {key} error: {e}")

        # ── SPLITS ───────────────────────────────────────────────────────────
        for key, info in st.session_state.sales_results.items():
            if do_split and n_splits > 1:
                info["splits"] = split_csv(info["tmp"], info["channel"],
                                           month, y, s, deal, int(n_splits), mls=False)

        for key, info in st.session_state.mls_results.items():
            if do_split and n_splits > 1:
                info["splits"] = split_csv(info["tmp"], info["channel"],
                                           month, y, s, deal, int(n_splits),
                                           mls=True, mls_type=info["mls_type"])

        # ── SALES ZIP ────────────────────────────────────────────────────────
        if st.session_state.sales_results:
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
                for key, info in st.session_state.sales_results.items():
                    folder = info["channel"]
                    if info["splits"]:
                        for wk_name, csv_bytes in info["splits"]:
                            zf.writestr(f"{folder}/{wk_name}.csv", csv_bytes)
                    else:
                        with open(info["tmp"], "rb") as f:
                            zf.writestr(f"{folder}/{info['campaign']}.csv", f.read())
            buf.seek(0)
            st.session_state.sales_zip_buffer = buf.getvalue()
            st.session_state.sales_zip_name   = f"Sales - {month} - {y} - {s} - {deal}.zip"

        # ── MLS SELLER ZIP ───────────────────────────────────────────────────
        mls_seller_out = {k: v for k, v in st.session_state.mls_results.items() if v.get("mls_type") == "SELLER"}
        if mls_seller_out:
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
                for key, info in mls_seller_out.items():
                    folder = info["channel"]
                    if info["splits"]:
                        for wk_name, csv_bytes in info["splits"]:
                            zf.writestr(f"{folder}/{wk_name}.csv", csv_bytes)
                    else:
                        with open(info["tmp"], "rb") as f:
                            zf.writestr(f"{folder}/{info['campaign']}.csv", f.read())
            buf.seek(0)
            st.session_state.mls_seller_zip_buffer = buf.getvalue()
            st.session_state.mls_seller_zip_name   = f"MLS Seller - {month} - {y} - {s} - {deal}.zip"

        # ── MLS AGENT ZIP ────────────────────────────────────────────────────
        mls_agent_out = {k: v for k, v in st.session_state.mls_results.items() if v.get("mls_type") == "AGENT"}
        if mls_agent_out:
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
                for key, info in mls_agent_out.items():
                    folder = info["channel"]
                    if info["splits"]:
                        for wk_name, csv_bytes in info["splits"]:
                            zf.writestr(f"{folder}/{wk_name}.csv", csv_bytes)
                    else:
                        with open(info["tmp"], "rb") as f:
                            zf.writestr(f"{folder}/{info['campaign']}.csv", f.read())
            buf.seek(0)
            st.session_state.mls_agent_zip_buffer = buf.getvalue()
            st.session_state.mls_agent_zip_name   = f"MLS Agent - {month} - {y} - {s} - {deal}.zip"

        # ── REPORT ───────────────────────────────────────────────────────────
        with st.spinner("Generating report... ⏳"):
            try:
                st.session_state.report_buffer = generate_report(
                    month, y, s, deal,
                    st.session_state.sales_results,
                    st.session_state.mls_results,
                    st.session_state.orig_phone_count,
                    st.session_state.orig_email_count,
                    int(n_splits), do_split
                )
            except Exception as e:
                st.error(f"❌ Report error: {e}")

        st.session_state.processed = True

    # ── RESULTS ──────────────────────────────────────────────────────────────
    if st.session_state.processed:
        st.divider()
        st.subheader("✅ Processing Complete!")

        st.info(f"📊 Total rows merged: **{st.session_state.total_merged:,}**")
        st.info(f"📞 Original phone numbers: **{st.session_state.orig_phone_count:,}**")
        st.info(f"📧 Original emails: **{st.session_state.orig_email_count:,}**")

        ICONS  = {"dialer": "📞", "sms": "💬", "email": "📧"}
        LABELS = {"dialer": "Dialer phones", "sms": "SMS phones", "email": "Emails"}

        if st.session_state.sales_results:
            st.markdown("**📋 Sales Process:**")
            for key, info in st.session_state.sales_results.items():
                ck = info["channel_key"]
                st.success(f"{ICONS[ck]} {LABELS[ck]}: **{info['rows']:,}**")

        mls_seller_out = {k: v for k, v in st.session_state.mls_results.items() if v.get("mls_type") == "SELLER"}
        mls_agent_out  = {k: v for k, v in st.session_state.mls_results.items() if v.get("mls_type") == "AGENT"}

        if mls_seller_out:
            st.markdown("**🏷️ MLS Seller List:**")
            for key, info in mls_seller_out.items():
                ck = info["channel_key"]
                st.success(f"{ICONS[ck]} {LABELS[ck]}: **{info['rows']:,}**")

        if mls_agent_out:
            st.markdown("**🧑‍💼 MLS Agent List:**")
            for key, info in mls_agent_out.items():
                ck = info["channel_key"]
                st.success(f"{ICONS[ck]} {LABELS[ck]}: **{info['rows']:,}**")

        st.divider()
        st.subheader("⬇️ Downloads")

        if st.session_state.sales_zip_buffer:
            st.download_button(
                label="⬇️ Download Sales Process (ZIP)",
                data=st.session_state.sales_zip_buffer,
                file_name=st.session_state.sales_zip_name,
                mime="application/zip",
                use_container_width=True,
                type="primary",
                key="dl_sales_zip"
            )

        if st.session_state.mls_seller_zip_buffer:
            st.download_button(
                label="⬇️ Download MLS Seller List (ZIP)",
                data=st.session_state.mls_seller_zip_buffer,
                file_name=st.session_state.mls_seller_zip_name,
                mime="application/zip",
                use_container_width=True,
                type="primary",
                key="dl_mls_seller_zip"
            )

        if st.session_state.mls_agent_zip_buffer:
            st.download_button(
                label="⬇️ Download MLS Agent List (ZIP)",
                data=st.session_state.mls_agent_zip_buffer,
                file_name=st.session_state.mls_agent_zip_name,
                mime="application/zip",
                use_container_width=True,
                type="primary",
                key="dl_mls_agent_zip"
            )

        if st.session_state.report_buffer:
            st.download_button(
                label="📄 Download Report (Word)",
                data=st.session_state.report_buffer,
                file_name=f"Report - {month} - {year} - {state.strip().upper()} - {deal}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
                key="dl_report"
            )

    if not ready and uploaded_files:
        if not state.strip():
            st.warning("⚠️ Please enter State.")
        if not year.strip():
            st.warning("⚠️ Please enter Year.")
        if not (sales_selected or mls_selected):
            st.warning("⚠️ Please select at least one output type.")

# ══════════════════════════════════════════════════════════════════════════
# PAGE: BUSINESS LEADS
# ══════════════════════════════════════════════════════════════════════════
def page_business_leads():
    st.title("🏢 Business Leads")
    st.divider()

    # STEP 1 — Upload
    st.subheader("Step 1 — Upload Files")
    uploaded_files = st.file_uploader(
        "📁 Upload one or more Excel or CSV files (.xlsx, .csv)",
        type=["xlsx", "csv"], accept_multiple_files=True, key="biz_upload"
    )
    if uploaded_files:
        st.success(f"✅ {len(uploaded_files)} file(s) uploaded")
    st.divider()

    # STEP 2 — Campaign Details
    st.subheader("Step 2 — Campaign Details")
    col1, col2 = st.columns(2)
    with col1:
        month = st.selectbox("📅 Month", MONTHS, key="biz_month")
        state = st.text_input("🗺️ State", placeholder="e.g. FL", key="biz_state")
    with col2:
        year = st.text_input("📆 Year", value="2026", key="biz_year")
        deal = st.text_input("🏷️ Niche / List Type", placeholder="e.g. HVAC CONTRACTORS", key="biz_deal")
    st.divider()

    # STEP 3 — Marketing Process
    st.subheader("Step 3 — Marketing Process")
    col_d, col_s, col_e = st.columns(3)
    with col_d:
        biz_dialer = st.checkbox("📞 Dialer", key="biz_dialer")
    with col_s:
        biz_sms    = st.checkbox("💬 SMS", key="biz_sms")
    with col_e:
        biz_email  = st.checkbox("📧 Email", key="biz_email")
    st.divider()

    # STEP 4 — Split
    st.subheader("Step 4 — Split (Optional)")
    do_split = st.checkbox("🔀 Split output into multiple files", key="biz_dosplit")
    n_splits = 1
    if do_split:
        n_splits = st.number_input("How many files to split into?",
                                   min_value=2, max_value=50, value=5, step=1, key="biz_nsplits")
    st.divider()

    # PROCESS
    biz_selected = biz_dialer or biz_sms or biz_email
    ready = bool(uploaded_files and state.strip() and year.strip() and deal.strip() and biz_selected)

    if st.button("⚙️ Process Files", use_container_width=True, type="primary",
                 disabled=not ready, key="biz_process_btn"):
        st.session_state.biz_processed = False
        st.session_state.biz_results   = {}

        with st.spinner("Merging files... ⏳"):
            try:
                total, orig_phones, orig_emails = merge_biz_to_disk(uploaded_files)
                st.session_state.biz_total_merged     = total
                st.session_state.biz_orig_phone_count = orig_phones
                st.session_state.biz_orig_email_count = orig_emails
            except Exception as e:
                st.error(f"❌ Merge error: {e}")
                st.stop()

        y = year.strip()
        s = state.strip().upper()
        d = deal.strip()

        # ── BUSINESS LEADS PROCESSING ────────────────────────────────────────
        BIZ_MAP = {
            "dialer": (biz_dialer, TMP_BIZ_DIALER, process_biz_dialer_file, "CC"),
            "sms":    (biz_sms,    TMP_BIZ_SMS,     process_biz_sms_file,    "SMS"),
            "email":  (biz_email,  TMP_BIZ_EMAIL,   process_biz_email_file,  "EMAIL"),
        }
        for key, (selected, out_tmp, fn, ch) in BIZ_MAP.items():
            if not selected:
                continue
            campaign = make_campaign(ch, month, y, s, d)
            with st.spinner(f"Processing Business {key.upper()}... ⏳"):
                try:
                    rows = fn(TMP_BIZ_MERGED, out_tmp, campaign)
                    st.session_state.biz_results[key] = {
                        "channel": ch, "channel_key": key, "campaign": campaign,
                        "rows": rows, "tmp": out_tmp, "splits": None,
                    }
                except Exception as e:
                    st.error(f"❌ Business {key} error: {e}")

        # ── SPLITS ───────────────────────────────────────────────────────────
        for key, info in st.session_state.biz_results.items():
            if do_split and n_splits > 1:
                info["splits"] = split_csv(info["tmp"], info["channel"],
                                           month, y, s, d, int(n_splits), mls=False)

        # ── ZIP ──────────────────────────────────────────────────────────────
        if st.session_state.biz_results:
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
                for key, info in st.session_state.biz_results.items():
                    folder = info["channel"]
                    if info["splits"]:
                        for wk_name, csv_bytes in info["splits"]:
                            zf.writestr(f"{folder}/{wk_name}.csv", csv_bytes)
                    else:
                        with open(info["tmp"], "rb") as f:
                            zf.writestr(f"{folder}/{info['campaign']}.csv", f.read())
            buf.seek(0)
            st.session_state.biz_zip_buffer = buf.getvalue()
            st.session_state.biz_zip_name   = f"Business Leads - {month} - {y} - {s} - {d}.zip"

        # ── REPORT ───────────────────────────────────────────────────────────
        with st.spinner("Generating report... ⏳"):
            try:
                st.session_state.biz_report_buffer = generate_report(
                    month, y, s, d,
                    st.session_state.biz_results,
                    {},  # no MLS section for Business Leads
                    st.session_state.biz_orig_phone_count,
                    st.session_state.biz_orig_email_count,
                    int(n_splits), do_split,
                    section_title="── BUSINESS LEADS ──"
                )
            except Exception as e:
                st.error(f"❌ Report error: {e}")

        st.session_state.biz_processed = True

    # ── RESULTS ──────────────────────────────────────────────────────────────
    if st.session_state.biz_processed:
        st.divider()
        st.subheader("✅ Processing Complete!")

        st.info(f"📊 Total rows merged: **{st.session_state.biz_total_merged:,}**")
        st.info(f"📞 Original phone numbers: **{st.session_state.biz_orig_phone_count:,}**")
        st.info(f"📧 Original emails: **{st.session_state.biz_orig_email_count:,}**")

        ICONS  = {"dialer": "📞", "sms": "💬", "email": "📧"}
        LABELS = {"dialer": "Dialer phones", "sms": "SMS phones", "email": "Emails"}

        if st.session_state.biz_results:
            st.markdown("**🏢 Business Leads:**")
            for key, info in st.session_state.biz_results.items():
                ck = info["channel_key"]
                st.success(f"{ICONS[ck]} {LABELS[ck]}: **{info['rows']:,}**")

        st.divider()
        st.subheader("⬇️ Downloads")

        if st.session_state.biz_zip_buffer:
            st.download_button(
                label="⬇️ Download Business Leads (ZIP)",
                data=st.session_state.biz_zip_buffer,
                file_name=st.session_state.biz_zip_name,
                mime="application/zip",
                use_container_width=True,
                type="primary",
                key="dl_biz_zip"
            )

        if st.session_state.biz_report_buffer:
            st.download_button(
                label="📄 Download Report (Word)",
                data=st.session_state.biz_report_buffer,
                file_name=f"Report - Business Leads - {month} - {year} - {state.strip().upper()} - {deal.strip()}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
                key="dl_biz_report"
            )

    if not ready and uploaded_files:
        if not state.strip():
            st.warning("⚠️ Please enter State.")
        if not year.strip():
            st.warning("⚠️ Please enter Year.")
        if not deal.strip():
            st.warning("⚠️ Please enter Niche / List Type.")
        if not biz_selected:
            st.warning("⚠️ Please select at least one channel (Dialer / SMS / Email).")

# ══════════════════════════════════════════════════════════════════════════
# APP ENTRY / NAVIGATION
# ══════════════════════════════════════════════════════════════════════════
st.set_page_config(page_title="List Cleaner", page_icon="📋", layout="centered")

NAV_PAGES = [("List Cleaner", "📋"), ("Business Leads", "🏢")]

with st.sidebar:
    st.markdown("## 📋 List Cleaner")
    st.divider()
    for name, icon in NAV_PAGES:
        is_active = st.session_state.current_page == name
        if st.button(f"{icon}  {name}", key=f"nav_{name}",
                     type="primary" if is_active else "secondary",
                     use_container_width=True):
            st.session_state.current_page = name
            st.rerun()

if st.session_state.current_page == "List Cleaner":
    page_sales_mls()
else:
    page_business_leads()
